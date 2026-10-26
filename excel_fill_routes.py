# -*- coding: utf-8 -*-
"""
Excel 占位符填充 Word —— 线上工具（接入 slowbuild 后端）
======================================================
路由:
  POST /api/excelfill/analyze       上传 Excel + Word 模板，返回列名/占位符预览
  POST /api/excelfill/run           执行填充，每行生成一个 docx 并打包 zip
  GET  /api/excelfill/download/<job> 下载结果 zip

前端页面: /excelfill.html (Vercel 托管)，BACKEND = api.slowbuild.top
"""

import os
import io
import re
import time
import uuid
import zipfile
import shutil

from flask import request, jsonify, send_file
from docx import Document

from excel_fill_core import (
    read_excel_headers_rows, fill_doc_bytes, safe_filename, PLACEHOLDER_RE,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
WORK_DIR = os.path.join(BASE_DIR, 'work', 'excelfill')
TEMP_DIR = os.path.join(WORK_DIR, 'temp')
os.makedirs(TEMP_DIR, exist_ok=True)

JOBS = {}
JOB_TTL = 3600  # 1 小时


def _cleanup_jobs():
    now = time.time()
    for job_id in list(JOBS.keys()):
        job = JOBS.get(job_id)
        if not job:
            continue
        if now - job['created'] > JOB_TTL:
            shutil.rmtree(job['dir'], ignore_errors=True)
            del JOBS[job_id]


def _extract_placeholders(doc):
    """提取文档中所有占位符（段落 + 表格），返回去重列表"""
    found = set()
    def scan_text(t):
        for m in PLACEHOLDER_RE.finditer(t or ''):
            found.add(m.group(0))
    for p in doc.paragraphs:
        scan_text(p.text)
    for table in doc.tables:
        for tr in table.rows:
            for cell in tr.cells:
                for p in cell.paragraphs:
                    scan_text(p.text)
    return sorted(found)


def register_excel_fill_routes(app):

    @app.route('/api/excelfill/analyze', methods=['POST'])
    def excelfill_analyze():
        _cleanup_jobs()

        excel = request.files.get('excel')
        template = request.files.get('template')
        if not excel:
            return jsonify({'error': '请上传 Excel 文件'}), 400
        if not template:
            return jsonify({'error': '请上传 Word 模板（.docx）'}), 400

        job_id = uuid.uuid4().hex[:12]
        job_dir = os.path.join(TEMP_DIR, job_id)
        os.makedirs(job_dir, exist_ok=True)

        excel_path = os.path.join(job_dir, 'data.xlsx')
        excel.save(excel_path)
        tpl_path = os.path.join(job_dir, 'template.docx')
        template.save(tpl_path)

        try:
            headers, rows = read_excel_headers_rows(excel_path)
        except Exception as e:
            shutil.rmtree(job_dir, ignore_errors=True)
            return jsonify({'error': '读取 Excel 失败: ' + str(e)}), 400

        try:
            doc = Document(tpl_path)
            placeholders = _extract_placeholders(doc)
        except Exception as e:
            shutil.rmtree(job_dir, ignore_errors=True)
            return jsonify({'error': '读取 Word 模板失败（请确认是 .docx 格式）: ' + str(e)}), 400

        JOBS[job_id] = {
            'created': time.time(),
            'dir': job_dir,
            'excel_path': excel_path,
            'tpl_path': tpl_path,
            'headers': headers,
            'rows_count': len(rows),
        }

        return jsonify({
            'job_id': job_id,
            'headers': headers,
            'rows_count': len(rows),
            'placeholders': placeholders,
            'guess_name_col': headers[0] if headers else None,
        })

    @app.route('/api/excelfill/run', methods=['POST'])
    def excelfill_run():
        data = request.get_json(silent=True) or {}
        job = JOBS.get(data.get('job_id'))
        if not job:
            return jsonify({'error': '会话已过期，请重新上传'}), 400

        name_col = data.get('name_col') or (job['headers'][0] if job['headers'] else '')
        name_col2 = data.get('name_col2') or ''

        _, rows = read_excel_headers_rows(job['excel_path'])

        out_dir = os.path.join(job['dir'], 'out')
        os.makedirs(out_dir, exist_ok=True)

        generated = 0
        all_filled = set()
        all_missing = set()

        for idx, row in enumerate(rows, start=1):
            data_bytes, filled, missing = fill_doc_bytes(job['tpl_path'], row)
            all_filled.update(filled)
            all_missing.update(missing)

            name1 = safe_filename(row.get(name_col, ''))
            if name_col2:
                name2 = safe_filename(row.get(name_col2, ''))
                fname = f"{idx:03d}_{name1}_{name2}.docx"
            else:
                fname = f"{idx:03d}_{name1}.docx"
            # 兜底防重名
            target = os.path.join(out_dir, fname)
            if os.path.exists(target):
                fname = f"{idx:03d}_{name1}_{uuid.uuid4().hex[:4]}.docx"
                target = os.path.join(out_dir, fname)
            with open(target, 'wb') as f:
                f.write(data_bytes)
            generated += 1

        # 打包 zip
        zip_path = os.path.join(job['dir'], 'result.zip')
        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
            for fname in sorted(os.listdir(out_dir)):
                zf.write(os.path.join(out_dir, fname), arcname=fname)

        job['zip_path'] = zip_path

        return jsonify({
            'count': generated,
            'filled': sorted(all_filled),
            'missing': sorted(all_missing),
        })

    @app.route('/api/excelfill/download/<job_id>', methods=['GET'])
    def excelfill_download(job_id):
        job = JOBS.get(job_id)
        if not job or not job.get('zip_path') or not os.path.exists(job['zip_path']):
            return jsonify({'error': '没有可下载的结果'}), 400
        return send_file(
            job['zip_path'], as_attachment=True, download_name='word填充结果.zip',
            mimetype='application/zip',
        )
