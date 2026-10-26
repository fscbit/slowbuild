# -*- coding: utf-8 -*-
"""
Excel/CSV 按列拆分 —— 核心逻辑
按某一列的不同取值，把表拆成多个文件（zip 下载）或一个文件多个 Sheet。
"""

import io
import zipfile
from table_utils import cell_to_str, rows_to_plain, write_xlsx_bytes


def split_by_column(headers, rows, col):
    """按 col 列的取值分组，返回 [(值, [行dict列表]), ...]（保持原顺序）"""
    groups = {}
    order = []
    for _, row in rows:
        v = cell_to_str(row.get(col))
        if v == '':
            v = '(空)'
        if v not in groups:
            groups[v] = []
            order.append(v)
        groups[v].append(row)
    return [(v, groups[v]) for v in order]


def safe_name(s, maxlen=50):
    """把分组值转成安全文件名"""
    s = (s or '').strip().replace('/', '_').replace('\\', '_').replace(':', '_') \
        .replace('*', '_').replace('?', '_').replace('"', '_').replace('<', '_') \
        .replace('>', '_').replace('|', '_')
    s = s.strip().strip('.') or '空'
    return s[:maxlen]


def build_split_zip(headers, groups, fmt='xlsx'):
    """按分组生成多个文件，打包成 zip 返回字节"""
    buf = io.BytesIO()
    used = set()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for value, rows in groups:
            base = safe_name(value)
            name = base
            k = 2
            while name in used:
                name = base + '_' + str(k)
                k += 1
            used.add(name)
            plain = rows_to_plain(headers, rows)
            if fmt == 'csv':
                from table_utils import write_csv_bytes
                zf.writestr(name + '.csv', write_csv_bytes(headers, plain))
            else:
                zf.writestr(name + '.xlsx', write_xlsx_bytes(headers, plain))
    buf.seek(0)
    return buf.getvalue()


def build_split_sheets(headers, groups):
    """按分组生成一个 xlsx，每个分组一个 Sheet"""
    if len(groups) > 100:
        raise RuntimeError('分组超过 100 个，请改用「多个文件」方式导出')
    from openpyxl.styles import Font, PatternFill, Alignment
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill('solid', fgColor='4472C4')
    head_font = Font(color='FFFFFF', bold=True)
    used = set()
    for value, rows in groups:
        base = safe_name(value, 28)
        name = base
        k = 2
        while name in used:
            name = base[:26] + '_' + str(k)
            k += 1
        used.add(name)
        ws = wb.create_sheet(title=name)
        ws.append(list(headers))
        for c in ws[1]:
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(horizontal='center', vertical='center')
        for row in rows_to_plain(headers, rows):
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
