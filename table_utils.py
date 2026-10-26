# -*- coding: utf-8 -*-
"""
通用表格读写工具（xlsx / csv）
供线上 Excel 工具（拆分/重命名/对比/合并/去重）共用，只用 openpyxl + 标准库。
"""

import io
import os
import csv
from datetime import datetime

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def cell_to_str(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0:
            return v.strftime('%Y-%m-%d')
        return v.strftime('%Y-%m-%d %H:%M')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _dedupe_headers(raw_headers):
    headers = []
    for h in raw_headers:
        h = str(h).strip() if h is not None else ''
        if not h:
            h = '列' + str(len(headers) + 1)
        base, k = h, 2
        while h in headers:
            h = base + '_' + str(k)
            k += 1
        headers.append(h)
    return headers


def _read_csv(path):
    raw = None
    for enc in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030'):
        try:
            with open(path, 'r', encoding=enc, newline='') as f:
                raw = f.read()
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if raw is None:
        with open(path, 'r', encoding='utf-8', errors='replace') as f:
            raw = f.read()

    lines = raw.splitlines()
    first_line = lines[0] if lines else ''
    sep = ','
    if first_line.count('\t') > first_line.count(','):
        sep = '\t'
    elif first_line.count(';') > first_line.count(','):
        sep = ';'

    reader = csv.reader(io.StringIO(raw), delimiter=sep)
    all_rows = list(reader)
    if not all_rows:
        return [], []
    headers = _dedupe_headers(all_rows[0])
    rows = []
    for i, r in enumerate(all_rows[1:], start=2):
        if not any(x is not None and str(x).strip() for x in r):
            continue
        d = {}
        for j, h in enumerate(headers):
            d[h] = r[j] if j < len(r) else None
        rows.append((i, d))
    return headers, rows


def read_table(path):
    """读取 xlsx 或 csv，返回 (headers, rows)；rows: [(原行号, {列名:值}), ...]"""
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        return _read_csv(path)
    if openpyxl is None:
        raise RuntimeError('缺少 openpyxl 依赖')
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = []
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            headers = _dedupe_headers(row)
            continue
        if not any(v is not None and str(v).strip() for v in row):
            continue
        d = {}
        for j, h in enumerate(headers):
            d[h] = row[j] if j < len(row) else None
        rows.append((i, d))
    wb.close()
    return headers, rows


def rows_to_plain(headers, rows):
    """rows [(行号, dict)] 或 [dict] → 纯值 list[list]（字符串）"""
    out = []
    for item in rows:
        if isinstance(item, tuple) and len(item) == 2 and isinstance(item[1], dict):
            r = item[1]
        elif isinstance(item, dict):
            r = item
        else:
            r = item
        out.append([cell_to_str(r.get(h)) for h in headers])
    return out


def write_xlsx_bytes(headers, rows_plain, highlight=None):
    """headers + rows_plain(list[list]) → xlsx 字节
    highlight: dict{(r,c): 'red'|'green'|'yellow'|'blue'}，r/c 从 1 开始（含表头行）
    """
    if openpyxl is None:
        raise RuntimeError('缺少 openpyxl 依赖')
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    head_fill = PatternFill('solid', fgColor='4472C4')
    head_font = Font(color='FFFFFF', bold=True)
    ws.append(list(headers))
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    fills = {
        'red': PatternFill('solid', fgColor='FFC7CE'),
        'green': PatternFill('solid', fgColor='C6EFCE'),
        'yellow': PatternFill('solid', fgColor='FFEB9C'),
        'blue': PatternFill('solid', fgColor='DDEBF7'),
    }
    for r, row in enumerate(rows_plain, start=2):
        ws.append(row)
        if highlight:
            for c in range(1, len(headers) + 1):
                col = highlight.get((r, c))
                if col and col in fills:
                    ws.cell(row=r, column=c).fill = fills[col]

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def write_csv_bytes(headers, rows_plain):
    """headers + rows_plain → CSV 字节（UTF-8 BOM，Excel 可直接打开）"""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for row in rows_plain:
        w.writerow(row)
    return ('\ufeff' + buf.getvalue()).encode('utf-8')


def resolve_column(sel, headers):
    """根据下拉值或列序号定位列名；返回列名或 None"""
    sel = (sel or '').strip()
    if not sel:
        return None
    if sel in headers:
        return sel
    if sel.isdigit():
        idx = int(sel) - 1
        if 0 <= idx < len(headers):
            return headers[idx]
    low = sel.lower()
    for h in headers:
        if h.lower() == low:
            return h
    return None
