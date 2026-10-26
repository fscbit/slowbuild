# -*- coding: utf-8 -*-
"""
Excel 占位符填充 Word —— 核心逻辑（纯函数，与 Web 解耦）

从方哥的 fill_word_from_excel_v3.1.py 提取占位符替换核心，通用化：
  · 读 Excel（第一行列名）→ 用 Word 模板（含 #列名# 占位符）→ 每行生成一个 docx
  · 占位符兼容半角 # 与全角 ＃
  · 列名匹配：优先精确匹配；否则把列名去掉首尾 #/＃ 后再匹配
  · 替换时先单 run 内替换（保留格式），再处理 Word 把占位符拆散到多个 run 的情况
"""

import io
import re
from datetime import datetime

from docx import Document
import openpyxl

# 占位符正则：兼容半角#/全角＃，允许中间有空格
PLACEHOLDER_RE = re.compile(r'[#＃][^#＃\r\n]{1,40}[#＃]')


def read_excel_headers_rows(path):
    """读取 Excel，返回 (headers, rows)。
    headers: 列名列表；rows: [{列名: 值}, ...]（值统一为字符串，空行为 ''）
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = []
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c).strip() if c is not None else '' for c in row]
            continue
        if not any(v is not None and str(v).strip() for v in row):
            continue
        d = {}
        for j, h in enumerate(headers):
            if not h:
                continue
            val = row[j] if j < len(row) else None
            if isinstance(val, datetime):
                val = val.strftime('%Y-%m-%d %H:%M:%S')
            elif val is None:
                val = ''
            else:
                val = str(val).strip()
            d[h] = val
        rows.append(d)
    wb.close()
    return headers, rows


def canon_ph(s):
    """把匹配到的占位符统一成纯列名：去掉首尾 #/＃、去内部空白、去全角空格
    例: '＃姓名＃' / '# 姓名 #' / '#姓名#' -> '姓名'
    """
    inner = s[1:-1]  # 去掉首尾定界符（半角#或全角＃各占1字符）
    inner = re.sub(r'\s+', '', inner)
    inner = inner.replace('　', '')
    return inner


def _find_value(key, data_row):
    """按纯列名 key 从数据行取值；精确优先，其次去 #/＃ 后比对。找不到返回 None"""
    if key in data_row:
        return data_row[key]
    for h, v in data_row.items():
        if h.strip('#＃').strip() == key:
            return v
    return None


def replace_in_paragraph(paragraph, data_row, filled_set, missing_set):
    """替换单个段落里的占位符（先单 run 内替换保格式，再处理跨 run 拆散）"""
    runs = paragraph.runs
    if not runs:
        return

    def _sub(m):
        key = canon_ph(m.group(0))
        val = _find_value(key, data_row)
        if val is None or val == '':
            missing_set.add(key)
            return m.group(0)  # 找不到对应列或值为空 → 原样保留，不乱填
        filled_set.add(key)
        return str(val)

    # 第1遍：占位符完整落在某个 run 内 → 直接替换，保留原格式
    for r in runs:
        if r.text and ('#' in r.text or '＃' in r.text):
            r.text = PLACEHOLDER_RE.sub(_sub, r.text)

    # 第2遍：占位符被 Word 拆散到多个 run → 合并整段文本重写
    full = ''.join(r.text or '' for r in runs)
    if not PLACEHOLDER_RE.search(full):
        return
    new = PLACEHOLDER_RE.sub(_sub, full)
    if new != full:
        runs[0].text = new
        for r in runs[1:]:
            r.text = ''


def fill_doc_bytes(template_path, data_row):
    """用数据行填充模板，返回 (bytes, filled_keys, missing_keys)"""
    doc = Document(template_path)
    filled = set()
    missing = set()

    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, data_row, filled, missing)
    for table in doc.tables:
        for tr in table.rows:
            for cell in tr.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, data_row, filled, missing)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue(), sorted(filled), sorted(missing)


def safe_filename(s):
    """清理非法文件名字符，返回安全文件名（空则返回占位）"""
    s = re.sub(r'[<>:"/\\|?*\r\n\t]', '', str(s or ''))
    return s.strip() or '未命名'
