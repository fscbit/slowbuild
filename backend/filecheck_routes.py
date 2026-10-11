# -*- coding: utf-8 -*-
"""
文件名与 Excel 核对 —— 线上工具（接入 slowbuild 后端）
======================================================
路由:
  POST /api/filecheck/analyze        上传 Excel + 文件(zip 或 文件夹多文件)
  POST /api/filecheck/check          执行核对
  GET  /api/filecheck/export/<job>   下载核对报告 xlsx

前端页面: /filecheck.html (Vercel 托管)，BACKEND = api.slowbuild.top
"""

import os
import io
import re
import time
import uuid
import zipfile
import shutil
from datetime import datetime

from flask import request, jsonify, send_file

from filecheck_core import (
    read_excel_rows, run_check, extract_field, guess_key_col, EXTRACT_METHODS,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
WORK_DIR = os.path.join(BASE_DIR, 'work', 'filecheck')
TEMP_DIR = os.path.join(WORK_DIR, 'temp')
os.makedirs(TEMP_DIR, exist_ok=True)

JOBS = {}
JOB_TTL = 3600  # 1 小时


def _safe_zip_name(name):
    """处理 Windows 打包 zip 的中文文件名乱码（cp437 -> gbk/utf-8）"""
    try:
        return name.encode('cp437').decode('gbk')
    except Exception:
        try:
            return name.encode('cp437').decode('utf-8')
        except Exception:
            return name


def _cleanup_jobs():
    now = time.time()
    for job_id in list(JOBS.keys()):
        job = JOBS.get(job_id)
        if not job:
            continue
        if now - job['created'] > JOB_TTL:
            shutil.rmtree(job['dir'], ignore_errors=True)
            del JOBS[job_id]


def _collect_files_from_zip(zip_path, dest_dir):
    names = []
    try:
        with zipfile.ZipFile(zip_path) as zf:
            for info in zf.infolist():
                name = _safe_zip_name(info.filename)
                safe = name.replace('\\', '/').lstrip('/')
                if '..' in safe.split('/'):
                    continue
                target = os.path.join(dest_dir, safe)
                if info.is_dir():
                    os.makedirs(target, exist_ok=True)
                    continue
                os.makedirs(os.path.dirname(target), exist_ok=True)
                with zf.open(info) as src, open(target, 'wb') as dst:
                    shutil.copyfileobj(src, dst)
                if os.path.isfile(target):
                    names.append(os.path.basename(name))
    except zipfile.BadZipFile:
        return None, '文件不是有效的 zip 压缩包'
    except Exception as e:
        return None, '解压失败: ' + str(e)
    return names, None


def _filter_exts(names, exts):
    if not exts:
        return names
    exts = {e.lower() if e.startswith('.') else '.' + e.lower() for e in exts}
    return [n for n in names if os.path.splitext(n)[1].lower() in exts]


def register_filecheck_routes(app):

    @app.route('/api/filecheck/analyze', methods=['POST'])
    def filecheck_analyze():
        _cleanup_jobs()

        excel = request.files.get('excel')
        if not excel:
            return jsonify({'error': '请上传 Excel 文件'}), 400

        job_id = uuid.uuid4().hex[:12]
        job_dir = os.path.join(TEMP_DIR, job_id)
        os.makedirs(job_dir, exist_ok=True)

        excel_path = os.path.join(job_dir, 'data.xlsx')
        excel.save(excel_path)

        file_names = []
        exts = [e.strip() for e in re.split(r'[,，;；\s]+', request.form.get('exts', '')) if e.strip()]

        zip_file = request.files.get('zip')
        if zip_file and zip_file.filename:
            zip_path = os.path.join(job_dir, 'files.zip')
            zip_file.save(zip_path)
            names, err = _collect_files_from_zip(zip_path, os.path.join(job_dir, 'files'))
            if err:
                shutil.rmtree(job_dir, ignore_errors=True)
                return jsonify({'error': err}), 400
            file_names = names
        else:
            uploads = [f for f in request.files.getlist('files') if f and f.filename]
            if not uploads:
                shutil.rmtree(job_dir, ignore_errors=True)
                return jsonify({'error': '请上传 zip 压缩包 或 选择文件夹'}), 400
            files_dir = os.path.join(job_dir, 'files')
            os.makedirs(files_dir, exist_ok=True)
            for f in uploads:
                rel = f.filename.replace('\\', '/').lstrip('/')
                if '..' in rel.split('/'):
                    continue
                target = os.path.join(files_dir, rel)
                os.makedirs(os.path.dirname(target), exist_ok=True)
                f.save(target)
                file_names.append(os.path.basename(f.filename))

        try:
            headers, rows = read_excel_rows(excel_path)
        except Exception as e:
            shutil.rmtree(job_dir, ignore_errors=True)
            return jsonify({'error': '读取 Excel 失败: ' + str(e)}), 400

        file_names = _filter_exts(sorted(set(file_names)), exts)

        JOBS[job_id] = {
            'created': time.time(),
            'dir': job_dir,
            'excel_path': excel_path,
            'headers': headers,
            'rows_count': len(rows),
            'file_names': file_names,
        }

        preview = []
        for name in file_names[:10]:
            preview.append({
                'name': name,
                'key': extract_field(name, '取前N个字', '9'),
            })

        return jsonify({
            'job_id': job_id,
            'headers': headers,
            'rows_count': len(rows),
            'files_count': len(file_names),
            'files': file_names,
            'preview': preview,
            'guess_key': guess_key_col(headers),
            'methods': EXTRACT_METHODS,
        })

    @app.route('/api/filecheck/check', methods=['POST'])
    def filecheck_check():
        data = request.get_json(silent=True) or {}
        job = JOBS.get(data.get('job_id'))
        if not job:
            return jsonify({'error': '会话已过期，请重新上传'}), 400

        headers = job['headers']
        _, rows = read_excel_rows(job['excel_path'])

        key_col = data.get('key_col')
        key_method = data.get('key_method')
        key_param = data.get('key_param', '')

        use_sub = bool(data.get('sub_enable'))
        sub_col = data.get('sub_col') if use_sub else None
        sub_method = data.get('sub_method', '')
        sub_param = data.get('sub_param', '')
        tol = data.get('tol', 0)

        if not key_col:
            return jsonify({'error': '请选择主核对列'}), 400
        if not key_method:
            return jsonify({'error': '请选择文件名切法'}), 400

        result = run_check(
            headers=headers, rows=rows, file_names=job['file_names'],
            key_col=key_col, key_method=key_method, key_param=key_param,
            sub_col=sub_col, sub_method=sub_method, sub_param=sub_param, tol=tol,
        )

        job['last_result'] = result
        job['last_config'] = {
            'key_col': key_col, 'key_method': key_method, 'key_param': key_param,
            'sub_col': sub_col, 'sub_method': sub_method, 'sub_param': sub_param,
            'tol': tol,
        }
        return jsonify({'result': result})

    @app.route('/api/filecheck/export/<job_id>', methods=['GET'])
    def filecheck_export(job_id):
        job = JOBS.get(job_id)
        if not job or 'last_result' not in job:
            return jsonify({'error': '没有可导出的结果'}), 400

        result = job['last_result']
        cfg = job['last_config']

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            return jsonify({'error': '缺少 openpyxl'}), 500

        wb = openpyxl.Workbook()
        head_fill = PatternFill('solid', fgColor='4472C4')
        head_font = Font(color='FFFFFF', bold=True)
        green = PatternFill('solid', fgColor='C6EFCE')
        yellow = PatternFill('solid', fgColor='FFEB9C')
        red = PatternFill('solid', fgColor='FFC7CE')
        blue = PatternFill('solid', fgColor='DDEBF7')

        ws1 = wb.active
        ws1.title = '逐行核对'
        ws1.append(['Excel行号', '主值', '副值', '状态', '匹配文件', '备注'])
        for c in ws1[1]:
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(horizontal='center', vertical='center')
        for r in result['rows']:
            ws1.append([r['rownum'], r['key'], r['sub'], r['status'], r['file'], r['note']])
            fill = {'ok': green, 'warn': yellow, 'err': red}.get(r['tag'])
            if fill:
                for c in ws1[ws1.max_row]:
                    c.fill = fill
        for idx, w in enumerate([10, 20, 20, 26, 46, 26], start=1):
            ws1.column_dimensions[get_column_letter(idx)].width = w

        ws2 = wb.create_sheet('多余文件')
        ws2.append(['文件名', '主字段提取', '副字段提取', '说明'])
        for c in ws2[1]:
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(horizontal='center', vertical='center')
        for ef in result['extra']:
            ws2.append([ef['name'], ef['key'], ef['sub'], ef['note']])
            for c in ws2[ws2.max_row]:
                c.fill = blue
        for idx, w in enumerate([50, 20, 20, 30], start=1):
            ws2.column_dimensions[get_column_letter(idx)].width = w

        ws3 = wb.create_sheet('汇总')
        s = result['stats']
        info = [
            ('主核对列', cfg['key_col']),
            ('主字段切法', cfg['key_method'] + ' / ' + str(cfg['key_param'])),
            ('副核对列', cfg['sub_col'] or '(未使用)'),
            ('副字段切法', (cfg['sub_method'] + ' / ' + str(cfg['sub_param'])) if cfg['sub_col'] else '(未使用)'),
            ('日期容差(天)', cfg['tol']),
            ('Excel 数据行数', s['excel_rows']),
            ('文件夹内文件数', s['files_total']),
            ('完全匹配', s['ok']),
            ('待核对', s['warn']),
            ('缺失文件', s['err']),
            ('多余文件', s['extra']),
            ('齐全率(%)', s['rate']),
        ]
        ws3.append(['项目', '数值'])
        for c in ws3[1]:
            c.fill = head_fill
            c.font = head_font
        for k, v in info:
            ws3.append([k, v])
        ws3.column_dimensions['A'].width = 20
        ws3.column_dimensions['B'].width = 60

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = '核对报告_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '.xlsx'
        return send_file(
            buf, as_attachment=True, download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
